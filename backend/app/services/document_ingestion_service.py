
import hashlib
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi import UploadFile
from pypdf import PdfReader

from app.schemas.document import DocumentCreate


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".pdf",
    ".docx",
}

MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB


class DocumentIngestionError(Exception):
    """Raised when an uploaded document cannot be processed."""


def get_file_extension(filename: str | None) -> str:
    """
    Return the lowercase extension of an uploaded filename.

    Example:
        policy.PDF -> .pdf
    """
    if not filename:
        return ""

    return Path(filename).suffix.lower()


def validate_file_extension(filename: str | None) -> str:
    """
    Validate that the uploaded file has a supported extension.

    Returns:
        The normalized lowercase extension.

    Raises:
        DocumentIngestionError: If the extension is unsupported.
    """
    extension = get_file_extension(filename)

    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentIngestionError(
            "Unsupported file type. Only TXT, Markdown, PDF, and DOCX "
            "files are allowed."
        )

    return extension


def validate_file_size(file_bytes: bytes) -> None:
    """
    Validate that the uploaded file is not empty and does not exceed
    the configured maximum size.
    """
    if not file_bytes:
        raise DocumentIngestionError("The uploaded file is empty.")

    if len(file_bytes) > MAX_FILE_SIZE:
        raise DocumentIngestionError(
            "The uploaded file exceeds the maximum allowed size of 5 MB."
        )


def extract_text_from_txt_or_markdown(file_bytes: bytes) -> str:
    """
    Extract UTF-8 text from a TXT or Markdown file.

    utf-8-sig supports regular UTF-8 and removes an optional BOM.
    """
    try:
        extracted_text = file_bytes.decode("utf-8-sig").strip()
    except UnicodeDecodeError as exc:
        raise DocumentIngestionError(
            "The uploaded text file must use UTF-8 encoding."
        ) from exc

    if not extracted_text:
        raise DocumentIngestionError(
            "No readable text was found in the uploaded file."
        )

    return extracted_text


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extract readable text from all pages of a text-based PDF.

    Image-only or scanned PDFs require OCR and are rejected for now.
    """
    try:
        reader = PdfReader(BytesIO(file_bytes))
    except Exception as exc:
        raise DocumentIngestionError(
            "The uploaded PDF file could not be read."
        ) from exc

    page_texts: list[str] = []

    for page in reader.pages:
        try:
            page_text = page.extract_text()
        except Exception:
            page_text = None

        if page_text and page_text.strip():
            page_texts.append(page_text.strip())

    extracted_text = "\n\n".join(page_texts).strip()

    if not extracted_text:
        raise DocumentIngestionError(
            "No readable text was found in the PDF. "
            "The file may be scanned or image-based."
        )

    return extracted_text


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extract readable text from DOCX paragraphs and tables.
    """
    try:
        document = Document(BytesIO(file_bytes))
    except Exception as exc:
        raise DocumentIngestionError(
            "The uploaded DOCX file could not be read."
        ) from exc

    extracted_parts: list[str] = []

    for paragraph in document.paragraphs:
        paragraph_text = paragraph.text.strip()

        if paragraph_text:
            extracted_parts.append(paragraph_text)

    for table in document.tables:
        for row in table.rows:
            cell_values = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if cell_values:
                extracted_parts.append(" | ".join(cell_values))

    extracted_text = "\n".join(extracted_parts).strip()

    if not extracted_text:
        raise DocumentIngestionError(
            "No readable text was found in the DOCX file."
        )

    return extracted_text


def extract_text_from_file(
    file_bytes: bytes,
    extension: str,
) -> str:
    """
    Select the correct text extractor based on the file extension.
    """
    if extension in {".txt", ".md"}:
        return extract_text_from_txt_or_markdown(file_bytes)

    if extension == ".pdf":
        return extract_text_from_pdf(file_bytes)

    if extension == ".docx":
        return extract_text_from_docx(file_bytes)

    raise DocumentIngestionError("Unsupported file type.")


async def read_and_validate_upload(
    file: UploadFile,
) -> tuple[bytes, str]:
    """
    Read an upload once and perform common validation.

    Returns:
        A tuple containing:
        - the uploaded file bytes
        - the normalized file extension
    """
    extension = validate_file_extension(file.filename)
    file_bytes = await file.read()

    validate_file_size(file_bytes)

    return file_bytes, extension


async def extract_text_from_upload(file: UploadFile) -> str:
    """
    Backward-compatible helper that validates an upload and returns
    only its extracted text.
    """
    file_bytes, extension = await read_and_validate_upload(file)

    return extract_text_from_file(
        file_bytes=file_bytes,
        extension=extension,
    )


def generate_title_from_filename(filename: str | None) -> str:
    """
    Generate a readable title from a filename.

    Example:
        remote_work-policy.pdf -> Remote Work Policy
    """
    if not filename:
        return "Untitled Document"

    stem = Path(filename).stem
    normalized_stem = stem.replace("_", " ").replace("-", " ")
    generated_title = " ".join(normalized_stem.split()).title()

    return generated_title or "Untitled Document"


def create_title_from_filename(filename: str | None) -> str:
    """
    Backward-compatible alias for generate_title_from_filename().
    """
    return generate_title_from_filename(filename)


def parse_tags(tags: str | list[str] | None) -> list[str]:
    """
    Normalize comma-separated tags or an existing tag list.

    Examples:
        "remote, policy, employees"
        -> ["remote", "policy", "employees"]

        ["remote", "policy"]
        -> ["remote", "policy"]
    """
    if not tags:
        return []

    raw_tags = tags if isinstance(tags, list) else tags.split(",")

    parsed_tags: list[str] = []
    seen_tags: set[str] = set()

    for tag in raw_tags:
        cleaned_tag = tag.strip()

        if not cleaned_tag:
            continue

        normalized_tag = cleaned_tag.casefold()

        if normalized_tag in seen_tags:
            continue

        seen_tags.add(normalized_tag)
        parsed_tags.append(cleaned_tag)

    return parsed_tags


def calculate_checksum(file_bytes: bytes) -> str:
    """
    Calculate the SHA-256 checksum of an uploaded file.
    """
    return hashlib.sha256(file_bytes).hexdigest()


async def process_uploaded_document(
    file: UploadFile,
    category: str,
    title: str | None = None,
    tags: str | list[str] | None = None,
    author: str = "Admin",
) -> DocumentCreate:
    """
    Validate an uploaded document, extract its text, calculate its
    checksum, and build the MongoDB document payload.
    """
    file_bytes, extension = await read_and_validate_upload(file)

    extracted_text = extract_text_from_file(
        file_bytes=file_bytes,
        extension=extension,
    )

    document_title = (
        title.strip()
        if title and title.strip()
        else generate_title_from_filename(file.filename)
    )

    normalized_category = category.strip()

    if not normalized_category:
        raise DocumentIngestionError(
            "The document category cannot be empty."
        )

    uploaded_at = datetime.now(timezone.utc)

    return DocumentCreate(
        title=document_title,
        category=normalized_category,
        content=extracted_text,
        tags=parse_tags(tags),
        author=author.strip() or "Admin",
        original_filename=file.filename,
        extension=extension,
        mime_type=file.content_type,
        file_size=len(file_bytes),
        checksum=calculate_checksum(file_bytes),
        uploaded_at=uploaded_at,
    )


async def ingest_document(
    file: UploadFile,
    category: str,
    title: str | None = None,
    tags: str | list[str] | None = None,
    author: str = "Admin",
) -> DocumentCreate:
    """
    Backward-compatible alias for process_uploaded_document().
    """
    return await process_uploaded_document(
        file=file,
        category=category,
        title=title,
        tags=tags,
        author=author,
    )
